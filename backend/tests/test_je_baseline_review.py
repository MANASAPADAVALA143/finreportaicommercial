"""JE history baseline, anomaly explanations, review queue and exports."""

from __future__ import annotations

import os
import sys
import types
from io import BytesIO
from types import SimpleNamespace

os.environ.setdefault("DATABASE_URL", "sqlite:///:memory:")

try:  # SciPy is unavailable on some dev machines; the engine itself is stubbed below.
    import scipy.stats  # noqa: F401
except Exception:
    _stats = types.ModuleType("scipy.stats")
    _stats.chisquare = lambda *a, **k: (0.0, 1.0)
    sys.modules.setdefault("scipy", types.ModuleType("scipy"))
    sys.modules["scipy.stats"] = _stats

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.api.routes import historical_analysis, history_router
from app.core import company_access
from app.core.database import get_db
from app.middleware.workspace import WorkspaceContext, validate_workspace
from app.models.history_models import (
    AccountBaseline,
    AccountMonthlyStat,
    JEAnomalyReview,
    JENarrative,
    JournalHistory,
)
from app.models.internal_audit import AuditEngagement, AuditFinding, AuditUniverse
from app.models.workspace import WorkspaceRole
from app.routers import internal_audit, internal_audit_je
from app.services.baseline_explainer import explain_against_baseline, severity_for_score

WS_A = "aaaaaaaa-0000-4000-8000-000000000001"
WS_B = "bbbbbbbb-0000-4000-8000-000000000002"
CO_A = "aaaaaaaa-1111-4000-8000-000000000001"
CO_A2 = "aaaaaaaa-1111-4000-8000-000000000003"
CO_B = "bbbbbbbb-1111-4000-8000-000000000002"
COMPANY_WORKSPACE = {CO_A: WS_A, CO_A2: WS_A, CO_B: WS_B}


class Env:
    def __init__(self):
        self.roles = {WS_A: "accountant"}

    def ctx(self) -> WorkspaceContext:
        role = self.roles.get(WS_A)
        return WorkspaceContext(
            workspace_id=WS_A,
            workspace=SimpleNamespace(id=WS_A),
            user=SimpleNamespace(id="user-1", email="auditor@example.com", role="user"),
            role=WorkspaceRole(role) if role else None,
        )


@pytest.fixture()
def env(monkeypatch):
    e = Env()
    monkeypatch.setattr(
        company_access, "_company_workspace_ids",
        lambda _db, cid: iter([COMPANY_WORKSPACE[cid]] if cid in COMPANY_WORKSPACE else []),
    )
    monkeypatch.setattr(company_access, "_workspace_role", lambda _db, _ctx, ws: e.roles.get(ws))
    return e


@pytest.fixture()
def session_factory():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    for model in (
        JournalHistory, AccountBaseline, AccountMonthlyStat, JEAnomalyReview, JENarrative,
        AuditUniverse, AuditEngagement, AuditFinding,
    ):
        model.__table__.create(engine)
    return sessionmaker(bind=engine)


def _app(env: Env | None, session_factory) -> FastAPI:
    app = FastAPI()
    app.include_router(history_router.router, prefix="/api/v2")
    app.include_router(historical_analysis.router, prefix="/api/v2")
    app.include_router(internal_audit.router)
    app.include_router(internal_audit_je.router)

    def _db():
        db = session_factory()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = _db
    if env is not None:
        app.dependency_overrides[validate_workspace] = env.ctx
    return app


@pytest.fixture()
def client(env, session_factory):
    return TestClient(_app(env, session_factory))


def _line(jid, date, account, amount, name="Professional Fees"):
    return {"journal_id": jid, "posting_date": date, "account": account, "account_name": name, "amount": amount}


HISTORY = [
    _line("J1", "2025-01-10", "4100", 100), _line("J2", "2025-01-20", "4100", 120),
    _line("J3", "2025-02-10", "4100", 110), _line("J4", "2025-02-20", "4100", 90),
    _line("J5", "2025-03-10", "4100", 130), _line("J6", "2025-03-20", "4100", 150),
    _line("J6", "2025-03-20", "2000", 150, "Accounts Payable"),
]


def _upload(client, entries=HISTORY, company=CO_A, rebuild=True):
    return client.post("/api/v2/history/upload", json={
        "company_id": company, "upload_month": "2025-03", "entries": entries, "rebuild": rebuild,
    })


# ── Login and isolation ──────────────────────────────────────────────────────

HISTORY_ROUTES = [
    ("POST", "/api/v2/history/upload"),
    ("GET", f"/api/v2/history/baseline-status?company_id={CO_A}"),
    ("GET", f"/api/v2/history/account-baseline?company_id={CO_A}"),
    ("DELETE", f"/api/v2/history/reset?company_id={CO_A}"),
    ("POST", "/api/v2/analyze-historical"),
    ("GET", f"/api/internal-audit/je-reviews?company_id={CO_A}"),
    ("PATCH", f"/api/internal-audit/je-reviews/1?company_id={CO_A}"),
    ("POST", f"/api/internal-audit/je-reviews/1/raise-finding?company_id={CO_A}"),
    ("GET", f"/api/internal-audit/je-reviews/export?company_id={CO_A}"),
]


@pytest.mark.parametrize("method,path", HISTORY_ROUTES, ids=[f"{m} {p.split('?')[0]}" for m, p in HISTORY_ROUTES])
def test_routes_require_login(env, session_factory, method, path):
    body = {"company_id": CO_A, "upload_month": "2025-03", "entries": []}
    res = TestClient(_app(None, session_factory)).request(method, path, json=body, headers={"X-Workspace-ID": WS_A})
    assert res.status_code == 401, res.text


def test_history_of_other_workspace_company_is_forbidden(client):
    assert _upload(client, company=CO_B).status_code == 403
    assert client.get(f"/api/v2/history/account-baseline?company_id={CO_B}").status_code == 403


def test_free_text_label_is_namespaced_to_workspace(client, session_factory):
    assert _upload(client, company="gnanova_demo").status_code == 200
    db = session_factory()
    assert {r.company_id for r in db.query(JournalHistory).all()} == {f"ws:{WS_A}:gnanova_demo"}
    status = client.get("/api/v2/history/baseline-status?company_id=gnanova_demo").json()
    assert status["company_id"] == "gnanova_demo" and status["months_loaded"] == 3


def test_viewer_cannot_upload_and_accountant_cannot_reset(client, env):
    env.roles[WS_A] = "viewer"
    assert _upload(client).status_code == 403
    env.roles[WS_A] = "accountant"
    assert _upload(client).status_code == 200
    assert client.delete(f"/api/v2/history/reset?company_id={CO_A}").status_code == 403
    env.roles[WS_A] = "finance_manager"
    assert client.delete(f"/api/v2/history/reset?company_id={CO_A}").json()["deleted_entries"] == 7
    assert client.get(f"/api/v2/history/account-baseline?company_id={CO_A}").json()["accounts_count"] == 0


# ── Baseline ─────────────────────────────────────────────────────────────────


def test_upload_builds_monthly_baseline_and_dedupes(client):
    res = _upload(client).json()
    assert res["saved"] == 7 and res["skipped_duplicates"] == 0
    assert res["months_in_baseline"] == 3 and res["accounts_covered"] == 2

    again = _upload(client).json()
    assert again["saved"] == 0 and again["skipped_duplicates"] == 7

    base = client.get(f"/api/v2/history/account-baseline?company_id={CO_A}").json()
    assert base["accounts_count"] == 2 and base["months_count"] == 3
    assert base["months"] == ["2025-01", "2025-02", "2025-03"]
    assert base["warnings"] and "at least 6 months" in base["warnings"][0]
    fees = next(a for a in base["accounts"] if a["account"] == "4100")
    assert fees["account_name"] == "Professional Fees"
    assert [m["total"] for m in fees["monthly"]] == [220, 200, 280]
    assert fees["avg_monthly_total"] == pytest.approx(233.33, abs=0.01)
    assert fees["alert_threshold"] == pytest.approx(233.33 + 2 * fees["std_monthly_total"], abs=0.02)
    assert fees["entry_count"] == 6 and fees["avg_entry_amount"] == pytest.approx(116.67, abs=0.01)


def test_chunked_upload_rebuilds_only_on_last_chunk(client):
    first = _upload(client, HISTORY[:4], rebuild=False).json()
    assert first["baseline_updated"] is False and "accounts_covered" not in first
    assert client.get(f"/api/v2/history/account-baseline?company_id={CO_A}").json()["accounts_count"] == 0
    last = _upload(client, HISTORY[4:], rebuild=True).json()
    assert last["months_in_baseline"] == 3 and last["total_entries_in_baseline"] == 7


# ── Explanation ──────────────────────────────────────────────────────────────


def _baseline(**kw):
    defaults = dict(mean_amount=115000, std_amount=20000, min_amount=80000, max_amount=145000,
                    months_loaded=18, total_entries=40)
    return SimpleNamespace(**{**defaults, **kw})


def test_explanation_for_outlier_matches_audit_wording():
    out = explain_against_baseline(
        account="4100", account_name="Professional Fees", amount=485000,
        posting_date="2025-03-15", baseline=_baseline(), currency="AED",
    )
    assert out["text"].startswith(
        "Account 4100 – Professional Fees posted AED 485,000 in March 2025 — 4.2× above the 18-month average of AED 115,000."
    )
    assert "Z-score: 18.5." in out["text"] and "Historical range: AED 80,000–AED 145,000." in out["text"]
    assert out["above_threshold"] is True and out["ratio"] == 4.22


def test_explanation_for_normal_amount_and_missing_baseline():
    normal = explain_against_baseline(
        account="4100", account_name=None, amount=118000, posting_date="2025-03-15", baseline=_baseline(),
    )
    assert "in line with the 18-month average" in normal["text"]
    assert "flag comes from other risk indicators" in normal["text"] and normal["above_threshold"] is False
    missing = explain_against_baseline(account="9999", account_name=None, amount=1, posting_date=None, baseline=None)
    assert "no historical baseline" in missing["text"] and missing["z_score"] is None


def test_severity_bands():
    assert [severity_for_score(s) for s in (90, 85.5, 85, 70, 69.9, 60.1, 60, 10)] == [
        "Critical", "Critical", "High", "High", "Medium", "Medium", None, None,
    ]


# ── Analysis → review queue ──────────────────────────────────────────────────


def _scored(jid, account, amount, score, level):
    return {
        "journal_id": jid, "account": account, "amount": amount, "user_id": "u1", "source": "ERP",
        "posting_date": "2025-04-15", "description": "", "entity": "",
        "composite": {"composite_score": score, "risk_level": level, "top_reasons": ["Amount spike"],
                      "layer_scores": {}},
        "layer_detail": {},
    }


@pytest.fixture()
def stub_engine(monkeypatch):
    scored = [
        _scored("N1", "4100", 900, 91, "CRITICAL"),
        _scored("N2", "4100", 130, 72, "HIGH"),
        _scored("N3", "2000", 150, 65, "MEDIUM"),
        _scored("N4", "4100", 110, 40, "LOW"),
    ]
    monkeypatch.setattr(historical_analysis._validator, "validate", lambda df, hist: {"ok": True})
    monkeypatch.setattr(historical_analysis._engine, "analyze", lambda df, hist, client_weights=None: {
        "entries": [dict(e, composite=dict(e["composite"])) for e in scored],
        "batch_stats": {"total": 4, "critical": 1, "high": 1, "medium": 1, "low": 1, "benford": {}},
    })
    prompts: list[dict] = []

    async def fake_batch(items, max_entries=20):
        prompts.extend(items)
        return {i["entry"]["journal_id"]: f"Observation for {i['entry']['journal_id']}." for i in items}

    monkeypatch.setattr(historical_analysis._narrative, "generate_batch", fake_batch)
    return prompts


def _analyze(client, company=CO_A):
    entries = [{"journal_id": f"N{i}", "posting_date": "2025-04-15", "account": "4100", "amount": 100} for i in range(1, 5)]
    return client.post("/api/v2/analyze-historical", json={"company_id": company, "entries": entries, "currency": "AED"})


def test_analysis_explains_flagged_entries_and_fills_review_queue(client, stub_engine):
    _upload(client)
    res = _analyze(client)
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["company_id"] == CO_A and body["review_queued"] == 3
    by_id = {e["journal_id"]: e for e in body["entries"]}
    assert "Account 4100 – Professional Fees posted AED 900 in April 2025" in by_id["N1"]["baseline_explanation"]["text"]
    assert by_id["N4"]["baseline_explanation"] is None
    assert by_id["N1"]["audit_narrative"] == "Observation for N1."
    assert {p["entry"]["journal_id"] for p in stub_engine} == {"N1", "N2", "N3"}
    assert "posted AED 900" in next(p for p in stub_engine if p["entry"]["journal_id"] == "N1")["entry"]["baseline_context"]

    queue = client.get(f"/api/internal-audit/je-reviews?company_id={CO_A}").json()
    assert [(r["journal_id"], r["severity"]) for r in queue["items"]] == [
        ("N1", "Critical"), ("N2", "High"), ("N3", "Medium"),
    ]
    assert queue["summary"]["by_severity"] == {"Critical": 1, "High": 1, "Medium": 1}
    assert queue["items"][0]["ai_observation"] == "Observation for N1."


def test_reanalysis_keeps_auditor_status_and_comment(client, stub_engine):
    _analyze(client)
    rid = client.get(f"/api/internal-audit/je-reviews?company_id={CO_A}").json()["items"][0]["id"]
    client.patch(f"/api/internal-audit/je-reviews/{rid}?company_id={CO_A}", json={"status": "Cleared", "auditor_comment": "Invoice seen"})
    _analyze(client)
    items = client.get(f"/api/internal-audit/je-reviews?company_id={CO_A}").json()["items"]
    assert len(items) == 3
    row = next(r for r in items if r["id"] == rid)
    assert row["status"] == "Cleared" and row["auditor_comment"] == "Invoice seen"


def test_read_only_member_can_analyze_without_writing_queue(client, env, stub_engine):
    env.roles[WS_A] = "viewer"
    res = _analyze(client)
    assert res.status_code == 200 and res.json()["review_queued"] == 0
    env.roles[WS_A] = "accountant"
    assert client.get(f"/api/internal-audit/je-reviews?company_id={CO_A}").json()["items"] == []


# ── Review actions ───────────────────────────────────────────────────────────


def _first_review(client, company=CO_A):
    return client.get(f"/api/internal-audit/je-reviews?company_id={company}").json()["items"][0]


def test_review_update_records_reviewer_and_validates(client, env, stub_engine):
    _analyze(client)
    rid = _first_review(client)["id"]
    url = f"/api/internal-audit/je-reviews/{rid}?company_id={CO_A}"
    res = client.patch(url, json={"status": "Under Review", "auditor_comment": "  Asked CFO  "})
    assert res.status_code == 200
    assert res.json()["reviewed_by"] == "auditor@example.com" and res.json()["auditor_comment"] == "Asked CFO"
    assert client.patch(url, json={"status": "Done"}).status_code == 422
    assert client.patch(url, json={"status": None}).status_code == 422
    env.roles[WS_A] = "viewer"
    assert client.patch(url, json={"status": "Cleared"}).status_code == 403


def test_reviews_are_isolated_between_companies(client, stub_engine):
    _analyze(client)
    rid = _first_review(client)["id"]
    assert client.get(f"/api/internal-audit/je-reviews?company_id={CO_A2}").json()["items"] == []
    assert client.patch(f"/api/internal-audit/je-reviews/{rid}?company_id={CO_A2}", json={"status": "Cleared"}).status_code == 404
    assert client.post(f"/api/internal-audit/je-reviews/{rid}/raise-finding?company_id={CO_A2}", json={}).status_code == 404
    assert client.get(f"/api/internal-audit/je-reviews?company_id={CO_B}").status_code == 403


def test_raise_as_finding_creates_linked_internal_audit_finding(client, stub_engine):
    _upload(client)
    _analyze(client)
    review = _first_review(client)
    url = f"/api/internal-audit/je-reviews/{review['id']}/raise-finding?company_id={CO_A}"

    other = client.post(f"/api/internal-audit/engagements?company_id={CO_A2}", json={"title": "Other"}).json()
    assert client.post(url, json={"engagement_id": other["id"]}).status_code == 422

    eng = client.post(f"/api/internal-audit/engagements?company_id={CO_A}", json={"title": "GL review"}).json()
    res = client.post(url, json={"engagement_id": eng["id"]})
    assert res.status_code == 201, res.text
    finding = res.json()["finding"]
    assert finding["risk_rating"] == "Critical" and finding["evidence_ref"] == "JE N1"
    assert finding["finding_ref"].startswith("F-") and "posted AED 900" in finding["description"]
    assert "Observation for N1." in finding["description"]
    assert res.json()["review"]["finding_id"] == finding["id"] and res.json()["review"]["status"] == "Under Review"

    assert client.post(url, json={}).status_code == 409
    engagements = client.get(f"/api/internal-audit/engagements?company_id={CO_A}").json()
    assert engagements[0]["findings_count"] == 1


# ── Exports ──────────────────────────────────────────────────────────────────


def test_excel_export_lists_filtered_reviews(client, stub_engine):
    from openpyxl import load_workbook

    _analyze(client)
    res = client.get(f"/api/internal-audit/je-reviews/export?company_id={CO_A}&format=xlsx&severity=Critical&company_name=Acme")
    assert res.status_code == 200
    assert res.headers["content-type"].startswith("application/vnd.openxmlformats")
    wb = load_workbook(BytesIO(res.content))
    rows = list(wb["JE Review"].iter_rows(values_only=True))
    assert rows[0][:3] == ("JE ID", "Posting date", "Account") and [r[0] for r in rows[1:]] == ["N1"]
    summary = dict(wb["Summary"].iter_rows(values_only=True))
    assert summary["Company"] == "Acme" and summary["Total flagged entries"] == 1


def test_pdf_export_and_period_validation(client, stub_engine):
    _analyze(client)
    base = f"/api/internal-audit/je-reviews/export?company_id={CO_A}"
    res = client.get(base + "&format=pdf&period_start=2025-04-01&period_end=2025-04-30")
    assert res.status_code == 200 and res.content.startswith(b"%PDF")
    assert "attachment" in res.headers["content-disposition"]
    assert client.get(base + "&period_start=April").status_code == 422
    empty = client.get(f"/api/internal-audit/je-reviews?company_id={CO_A}&period_start=2025-05-01").json()
    assert empty["items"] == []
